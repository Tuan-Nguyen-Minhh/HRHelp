from datetime import date
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest

from src.agents.case_graph import build_case_graph
from src.agents.handoff import CaseHandoff, CitedPassage
from src.core.config import settings
from src.features.answer.model import answers
from src.features.answer.model.agent import build_recommend_agent
from src.features.answer.model.assessment import (
    ApplicableRule,
    Assessment,
    CaseSummary,
    DownstreamHandoff,
    EvidenceRef,
    RiskAssessment,
    RuleCondition,
    SanctionAssessment,
)
from src.features.answer.model.case_tools import disciplinary_minutes_document, export_word_document
from src.features.answer.model.decide import DecisionInput, decide, should_synthesize
from src.features.answer.model.draft import AnswerInput, draft_answer
from src.features.answer.model.layers import LayeredBasis, split_passages
from src.features.answer.model.standard import applied_standard
from src.features.answer.model.synthesis import PROMPT_VERSION, SYSTEM_PROMPT, guard, synthesize
from src.features.answer.schema.render import HEADINGS
from src.features.assessment.schemas import (
    AssessmentResult,
    AssessmentSection,
    CaseInformation,
    EmployeeInformation,
    EvidenceCard,
    HandlingProposal,
    IncidentSummary,
    RiskAndGaps,
    Sections,
)
from src.features.cite.model.passage import Passage
from src.features.cite.model.query import SearchQuery
from src.features.clarify.model.assess import StoredFact
from src.features.clarify.model.catalog import ABSENCE
from src.models.cases import Case

_READY = {
    "contract_type": "xác định thời hạn",
    "absence_days": "5 ngày trong 30 ngày",
    "notice_form": "không báo",
    "minutes_made": "true",
    "explanation_requested": "true",
}


def _passage(policy_type: str, code: str, name: str, content: str) -> Passage:
    return Passage(
        chunk_id=uuid4(),
        policy_version_id=uuid4(),
        content=content,
        article="125",
        clause="1",
        point=None,
        score=0.8,
        policy_type=policy_type,
        policy_code=code,
        policy_name=name,
    )


def _layered(*passages: Passage) -> LayeredBasis:
    return split_passages(list(passages))


_ART_124 = "1. Khiển trách.\n2. Kéo dài thời hạn nâng lương không quá 06 tháng.\n3. Cách chức.\n4. Sa thải."
_ART_125 = (
    "Hình thức xử lý kỷ luật sa thải được người sử dụng lao động áp dụng trong trường hợp sau đây:\n"
    "4. Người lao động tự ý bỏ việc 05 ngày cộng dồn trong thời hạn 30 ngày "
    "hoặc 20 ngày cộng dồn trong thời hạn 365 ngày tính từ ngày đầu tiên tự ý bỏ việc "
    "mà không có lý do chính đáng."
)
_ABSENCE = "nghỉ không phép tự ý bỏ việc"


def _decide(
    layered: LayeredBasis,
    facts: dict[str, str] | None = None,
    description: str = "",
    assessment: Assessment | None = None,
):
    return decide(
        DecisionInput(family=ABSENCE, facts=facts or _READY, layered=layered, description=description),
        assessment,
    )


def _supported(level: str, law: Passage, rule: Passage, quote_law: str | None = None, quote_rule: str | None = None):
    raw = Assessment(
        case_summary=CaseSummary(violation="nghỉ không phép", facts=dict(_READY)),
        sanction_assessment=SanctionAssessment(
            status="SUPPORTED",
            recommended_sanction=level,
            reason="Đoạn đã lấy nêu mức này.",
            evidence=[
                EvidenceRef(chunk_id=str(law.chunk_id), quote=quote_law or law.content),
                EvidenceRef(chunk_id=str(rule.chunk_id), quote=quote_rule or rule.content),
            ],
        ),
        risk_assessment=RiskAssessment(level="LOW", reasons=["Đủ trích dẫn."]),
        applicable_rules=[ApplicableRule(
            rule_id=str(item.chunk_id), document_type=item.policy_type, description="Điều khoản thử nghiệm",
            supports_sanction=level, evidence=[EvidenceRef(chunk_id=str(item.chunk_id), quote=quote)],
            conditions=[RuleCondition(fact="absence_days", fact_value=_READY["absence_days"],
                         required="Điều kiện trong fixture", result="SATISFIED",
                         evidence=[EvidenceRef(chunk_id=str(item.chunk_id), quote=quote)])],
        ) for item, quote in ((law, quote_law or law.content), (rule, quote_rule or rule.content))],
        downstream_handoff=DownstreamHandoff(can_recommend=True),
    )
    return guard(raw, [law, rule])


def test_disciplinary_minutes_document_is_returned_even_without_a_sanction():
    case = Case(
        id=uuid4(),
        created_by=uuid4(),
        case_code="CASE-001",
        employee_code="NV001",
        employee_name="Nguyễn Văn A",
        description="Phòng ban: SOFTWARE\nChức danh: DEVELOPER\nVi phạm: Người lao động nghỉ không phép 5 ngày.",
        incident_date=date(2026, 3, 1),
        status="ANALYZED",
        family=ABSENCE,
    )
    document = disciplinary_minutes_document(case, None)

    assert document.title == "BIÊN BẢN XỬ LÝ VI PHẠM KỶ LUẬT"
    assert "Ông/Bà: Nguyễn Văn A" in document.content
    assert "Chức vụ/Chức danh: DEVELOPER" in document.content
    assert "Phòng ban/Bộ phận: SOFTWARE" in document.content
    assert "Đối chiếu quy định" in document.content
    assert "Bằng chứng, tang vật" not in document.content
    assert "Cuộc họp kết thúc" not in document.content
    content_section = document.content.split("- Diễn biến sự việc:", maxsplit=1)[1]
    content_section = content_section.split("- Đối chiếu quy định:", maxsplit=1)[0]
    assert "Phòng ban:" not in content_section
    assert "Chức danh:" not in content_section
    assert "Hình thức xử phạt của Ban lãnh đạo: ......................................................................................." in document.content
    assert "Đại diện Ban lãnh đạo" in document.content


def test_disciplinary_minutes_exports_word_form(monkeypatch, tmp_path):
    from docx import Document

    from src.features.answer.model import case_tools

    monkeypatch.setattr(case_tools, "_EXPORT_DIR", tmp_path)
    case = Case(
        id=uuid4(),
        created_by=uuid4(),
        case_code="CASE-002",
        employee_code="NV002",
        employee_name="Trần Thị B",
        description="Đi muộn nhiều lần.",
        incident_date=date(2026, 4, 2),
        status="ANALYZED",
        family="LATE",
    )
    path = export_word_document(disciplinary_minutes_document(case, None), "bien-ban-test")

    doc = Document(path)
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    table_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert "BIÊN BẢN XỬ LÝ VI PHẠM KỶ LUẬT" in text
    assert "Trần Thị B" in text
    assert "CÔNG TY" in table_text
    assert "Người bị lập biên bản" in table_text


def test_disciplinary_minutes_cleans_raw_report_from_incident_narrative():
    case = Case(
        id=uuid4(),
        created_by=uuid4(),
        case_code="CASE-003",
        employee_code="NV003",
        employee_name="Nguyễn Văn C",
        description="Tự ý nghỉ không phép trong 5 ngày liên tục.",
        incident_date=date(2026, 10, 20),
        status="ANALYZED",
        family=ABSENCE,
    )
    sections = Sections(
        case_information=CaseInformation(case_code="CASE-003", incident_date=date(2026, 10, 20), case_status="ANALYZED"),
        employee_information=EmployeeInformation(employee_code="NV003", employee_name="Nguyễn Văn C"),
        incident_summary=IncidentSummary(
            narrative=(
                "BIEN BAN ĐANH GIA VU VIEC I. THONG TIN VU VIEC - Nhân viên: Nguyễn Văn C. "
                "II. TÌNH TIẾT - Chunk_id: 5d8458e0-04af-44fa-970c-157922461604."
            )
        ),
        confirmed_facts=[],
        evidence_reviewed=[],
        assessment=AssessmentSection(),
        assessment_result=AssessmentResult(status="INSUFFICIENT_EVIDENCE", summary=""),
        handling_proposal=HandlingProposal(),
        risk_and_gaps=RiskAndGaps(),
        next_steps=[],
        citations=[],
    )
    document = disciplinary_minutes_document(case, None, sections)

    assert "- Diễn biến sự việc:\n  - Tự ý nghỉ không phép trong 5 ngày liên tục." in document.content
    assert "Chunk_id" not in document.content
    assert "BIEN BAN ĐANH GIA" not in document.content


def test_disciplinary_minutes_orders_rule_decree_law_in_basis():
    case = Case(
        id=uuid4(),
        created_by=uuid4(),
        case_code="CASE-004",
        employee_code="NV004",
        employee_name="Nguyễn Văn D",
        description="Đi muộn nhiều lần.",
        incident_date=date(2026, 10, 21),
        status="ANALYZED",
        family="LATE",
    )
    sections = Sections(
        case_information=CaseInformation(case_code="CASE-004", incident_date=date(2026, 10, 21), case_status="ANALYZED"),
        employee_information=EmployeeInformation(employee_code="NV004", employee_name="Nguyễn Văn D"),
        incident_summary=IncidentSummary(narrative="- Đi muộn nhiều lần."),
        confirmed_facts=[],
        evidence_reviewed=[],
        assessment=AssessmentSection(),
        assessment_result=AssessmentResult(status="INSUFFICIENT_EVIDENCE", summary="Chưa đủ căn cứ."),
        handling_proposal=HandlingProposal(),
        risk_and_gaps=RiskAndGaps(),
        next_steps=[],
        citations=[
            EvidenceCard(evidence_id="EV-001", policy_type="LAW", title="Bộ luật Lao động", citation_label="Bộ luật Lao động — Điều 34", quote="", layer="legal_basis"),
            EvidenceCard(evidence_id="EV-002", policy_type="DECREE", title="Nghị định 283", citation_label="Nghị định 283 — Điều 17", quote="", layer="legal_basis"),
            EvidenceCard(evidence_id="EV-003", policy_type="INTERNAL_RULE", title="Nội quy lao động", citation_label="Nội quy lao động — Điều 11", quote="", layer="internal_rule"),
        ],
    )

    document = disciplinary_minutes_document(case, None, sections)
    basis = document.content.split("- Đối chiếu quy định: ", maxsplit=1)[1].splitlines()[0]

    assert basis.index("Nội quy lao động") < basis.index("Nghị định 283") < basis.index("Bộ luật Lao động")


def test_python_does_not_infer_a_window_or_a_level_from_the_passage_text():
    layered = _layered(
        _passage("LAW", "BLLD", "Bộ luật Lao động", _ART_125),
        _passage("INTERNAL_RULE", "NQ", "Nội quy", "Người lao động tự ý bỏ việc không phép thì bị sa thải."),
    )
    assert _decide(layered, description=_ABSENCE).sanction_type is None
    assert _decide(layered, {**_READY, "absence_days": "5 ngày"}, _ABSENCE).sanction_type is None
    assert _decide(layered, {**_READY, "absence_days": "5 ngày trong 30 ngày"}, _ABSENCE).sanction_type is None


def test_no_corpus_does_not_lock_a_sanction():
    decision = _decide(_layered())
    assert decision.sanction_type is None
    assert decision.risk_level == "CRITICAL"


def test_decree_283_alone_cannot_set_a_company_sanction():
    penalty = _passage(
        "DECREE", "283/2026/ND-CP", "Nghị định xử phạt vi phạm hành chính", "Phạt tiền người sử dụng lao động."
    )
    decision = _decide(_layered(penalty))
    assert decision.sanction_type is None
    assert decision.risk_level == "HIGH"
    assert decision.sanction_type != "TERMINATION"


def test_missing_internal_rule_keeps_the_level_open():
    law = _passage("LAW", "BLLD-2019", "Bộ luật Lao động", "Người lao động có thể bị sa thải.")
    decision = _decide(_layered(law))
    assert decision.sanction_type is None
    assert decision.risk_level == "HIGH"
    assert decision.missing_internal_rule is True


def test_internal_rule_without_law_does_not_lock_a_sanction():
    rule = _passage("INTERNAL_RULE", "NQ", "Nội quy", "Nghỉ không phép thì bị sa thải.")
    decision = _decide(_layered(rule))
    assert decision.sanction_type is None
    assert decision.risk_level == "HIGH"


def test_page_index_headings_do_not_lock_a_sanction():
    layered = _layered(
        _passage("LAW", "BLLD-2019", "Bộ luật Lao động", "Sa thải."),
        _passage("INTERNAL_RULE", "NQ", "Nội quy", "Nghỉ không phép thì bị sa thải."),
    )
    decision = decide(DecisionInput(family=ABSENCE, facts=_READY, layered=layered, retrieval="page_index"))
    assert decision.sanction_type is None
    assert decision.risk_level == "HIGH"


def test_incomplete_procedure_blocks_dismissal():
    layered = _layered(
        _passage("LAW", "BLLD-2019", "Bộ luật Lao động", "Sa thải."),
        _passage("INTERNAL_RULE", "NQ", "Nội quy", "Tự ý nghỉ không phép thì bị sa thải."),
    )
    facts = {**_READY, "minutes_made": "false"}
    decision = _decide(layered, facts)
    assert decision.sanction_type is None
    assert decision.risk_level == "HIGH"


def test_dismissal_is_not_termination():
    law = _passage("LAW", "BLLD-2019", "Bộ luật Lao động", _ART_125)
    rule = _passage("INTERNAL_RULE", "NQ", "Nội quy", "Người lao động tự ý bỏ việc không phép thì bị sa thải.")
    layered = _layered(law, rule)
    assessment = _supported(
        "DISMISSAL",
        law,
        rule,
        quote_law=_ART_125,
    )
    decision = _decide(layered, description=_ABSENCE, assessment=assessment)
    assert decision.sanction_type == "DISMISSAL"
    assert "TERMINATION" in decision.rationale
    prose = draft_answer(
        AnswerInput(
            family=ABSENCE,
            facts=_READY,
            fact_lines=["Loại hợp đồng: xác định thời hạn"],
            employee_name="An",
            incident_date=date(2026, 3, 1),
            layered=layered,
            precedent_note="Không có đoạn tiền lệ trong lượt đối chiếu này.",
            description=_ABSENCE,
        ),
        assessment,
    ).prose
    assert "DISMISSAL" in prose
    assert "Không đề xuất chấm dứt hợp đồng (TERMINATION)." in prose
    for heading in HEADINGS:
        assert heading in prose


def test_termination_wording_in_an_internal_rule_is_not_selected():
    layered = _layered(
        _passage("LAW", "BLLD-2019", "Bộ luật Lao động", "Chấm dứt hợp đồng lao động."),
        _passage("INTERNAL_RULE", "NQ", "Nội quy", "Công ty được chấm dứt hợp đồng."),
    )
    facts = {**_READY, "absence_days": "1 ngày trong 30 ngày"}
    decision = _decide(layered, facts)
    assert decision.sanction_type is None


def test_conflicting_sources_collapse_to_one_open_standard():
    layered = _layered(
        _passage("LAW", "BLLD-2019", "Bộ luật Lao động", "Có thể khiển trách hoặc cảnh cáo."),
        _passage("INTERNAL_RULE", "NQ", "Nội quy", "Vi phạm thì khiển trách hoặc cảnh cáo."),
        _passage("DECREE", "283/2026/ND-CP", "Nghị định xử phạt", "Phạt tiền và sa thải."),
        _passage("PRECEDENT", "QD-1", "Vụ cũ", "Vụ trước đã sa thải."),
    )
    facts = {**_READY, "absence_days": "1 ngày trong 30 ngày"}
    decision = _decide(layered, facts)
    standard = applied_standard(decision.sanction_type, decision.risk_level, decision.process_steps)
    prose = draft_answer(
        AnswerInput(
            family=ABSENCE,
            facts=facts,
            fact_lines=["Loại hợp đồng: xác định thời hạn"],
            employee_name="An",
            incident_date=date(2026, 3, 1),
            layered=layered,
            precedent_note="Có 1 đoạn tiền lệ. Tiền lệ chỉ để đối chiếu, không tự nâng mức kỷ luật.",
        )
    ).prose
    assert decision.sanction_type is None
    assert standard.aligned is False
    assert standard.sanction_type is None
    assert len(standard.process_steps) == 4
    assert [item.effect for item in standard.layers] == ["basis", "open", "compare", "excluded"]
    assert "Chuẩn áp dụng" in prose
    assert "Nghị định xử phạt không vào mức kỷ luật." in prose
    assert standard.sanction_type != "TERMINATION"


def test_a_grounded_quote_can_lock_warning_and_an_invented_quote_cannot():
    law = _passage("LAW", "BLLD-2019", "Bộ luật Lao động", _ART_124)
    rule = _passage("INTERNAL_RULE", "NQ", "Nội quy", "Đi muộn lần đầu thì khiển trách.")
    layered = _layered(law, rule)
    grounded = _supported("WARNING", law, rule, quote_law="1. Khiển trách.")
    invented = _supported("WARNING", law, rule, quote_rule="Đi muộn thì sa thải.")
    assert grounded.downstream_handoff.can_recommend is True
    assert _decide(layered, assessment=grounded).sanction_type == "WARNING"
    assert invented.downstream_handoff.can_recommend is False
    assert _decide(layered, assessment=invented).sanction_type is None


def test_multiple_supported_levels_are_not_reduced_to_one():
    law = _passage("LAW", "BLLD", "Bộ luật Lao động", _ART_124)
    rule = _passage("INTERNAL_RULE", "NQ", "Nội quy", "Đi muộn lần đầu thì khiển trách hoặc cách chức.")
    raw = Assessment(
        case_summary=CaseSummary(violation="đi muộn", facts={}),
        sanction_assessment=SanctionAssessment(
            status="MULTIPLE_SUPPORTED",
            recommended_sanction="DEMOTION",
            reason="Nội quy nêu hai mức.",
            evidence=[
                EvidenceRef(chunk_id=str(law.chunk_id), quote="1. Khiển trách."),
                EvidenceRef(chunk_id=str(rule.chunk_id), quote=rule.content),
            ],
        ),
        risk_assessment=RiskAssessment(level="MEDIUM", reasons=["Hai mức."]),
    )
    accepted = guard(raw, [law, rule])
    decision = _decide(_layered(law, rule), assessment=accepted)
    assert accepted.sanction_assessment.recommended_sanction is None
    assert decision.sanction_type is None
    assert decision.risk_level == "MEDIUM"


def test_an_administrative_penalty_cannot_ground_a_company_sanction():
    penalty = _passage("DECREE", "283/2026/ND-CP", "Nghị định xử phạt", "Phạt tiền và sa thải.")
    precedent = _passage("PRECEDENT", "QD-1", "Vụ cũ", "Vụ trước đã sa thải.")
    rule = _passage("INTERNAL_RULE", "NQ", "Nội quy", "Đi muộn lần đầu thì khiển trách.")
    raw = Assessment(
        case_summary=CaseSummary(violation="đi muộn", facts={}),
        sanction_assessment=SanctionAssessment(
            status="SUPPORTED",
            recommended_sanction="DISMISSAL",
            reason="Nghị định xử phạt nói sa thải.",
            evidence=[
                EvidenceRef(chunk_id=str(penalty.chunk_id), quote="Phạt tiền và sa thải."),
                EvidenceRef(chunk_id=str(precedent.chunk_id), quote="Vụ trước đã sa thải."),
                EvidenceRef(chunk_id=str(rule.chunk_id), quote=rule.content),
            ],
        ),
        risk_assessment=RiskAssessment(level="LOW", reasons=["Sai nguồn."]),
        downstream_handoff=DownstreamHandoff(can_recommend=True),
    )
    accepted = guard(raw, [penalty, precedent, rule])
    assert accepted.downstream_handoff.can_recommend is False
    assert _decide(_layered(penalty, precedent, rule), assessment=accepted).sanction_type is None


def test_need_more_facts_clears_a_filled_sanction():
    law = _passage("LAW", "BLLD", "Bộ luật Lao động", _ART_125)
    rule = _passage("INTERNAL_RULE", "NQ", "Nội quy", "Người lao động tự ý bỏ việc không phép thì bị sa thải.")
    raw = Assessment(
        case_summary=CaseSummary(violation="nghỉ không phép", facts={"absence_days": "5 ngày"}),
        sanction_assessment=SanctionAssessment(
            status="NEED_MORE_FACTS",
            recommended_sanction="DISMISSAL",
            reason="Fact chỉ có 5 ngày, đoạn luật yêu cầu mốc 30 ngày.",
            evidence=[
                EvidenceRef(chunk_id=str(law.chunk_id), quote="05 ngày cộng dồn trong thời hạn 30 ngày"),
                EvidenceRef(chunk_id=str(rule.chunk_id), quote=rule.content),
            ],
        ),
        risk_assessment=RiskAssessment(level="LOW", reasons=["Thiếu mốc."]),
        downstream_handoff=DownstreamHandoff(can_recommend=True),
    )
    accepted = guard(raw, [law, rule])
    facts = {**_READY, "absence_days": "5 ngày"}
    decision = _decide(_layered(law, rule), facts, assessment=accepted)
    assert accepted.sanction_assessment.recommended_sanction is None
    assert accepted.sanction_assessment.status == "NEED_MORE_FACTS"
    assert decision.sanction_type is None
    assert decision.risk_level == "HIGH"


def test_synthesis_prompt_returns_a_level_when_the_passages_state_one():
    assert PROMPT_VERSION == "evidence-synthesis-v6-general-rubric"
    assert "status = SUPPORTED" in SYSTEM_PROMPT
    assert "BIÊN BẢN ĐÁNH GIÁ VỤ VIỆC" in SYSTEM_PROMPT
    assert "Chưa đủ căn cứ retrieval để xác định mức xử lý." in SYSTEM_PROMPT
    assert "NEED_MORE_FACTS" in SYSTEM_PROMPT
    assert "Không chọn hình thức nặng nhất" in SYSTEM_PROMPT
    assert "TERMINATION" in SYSTEM_PROMPT
    assert "can_create_legal_ticket = false" in SYSTEM_PROMPT
    assert "SEND_TO_LEGAL" in SYSTEM_PROMPT
    assert "UNKNOWN" in SYSTEM_PROMPT


def test_internal_rule_level_stays_in_the_report_and_is_not_locked():
    rule = _passage("INTERNAL_RULE", "NQ", "Nội quy", "Tự ý bỏ việc 01 ngày thì khiển trách.")
    raw = Assessment(
        case_summary=CaseSummary(violation="nghỉ không phép", facts={"absence_days": "1 ngày"}),
        sanction_assessment=SanctionAssessment(
            status="SUPPORTED",
            recommended_sanction="WARNING",
            reason="Nội quy nêu khiển trách cho 01 ngày.",
            evidence=[EvidenceRef(chunk_id=str(rule.chunk_id), quote=rule.content)],
        ),
        risk_assessment=RiskAssessment(level="LOW", reasons=["Theo nội quy."]),
        report="BIÊN BẢN ĐÁNH GIÁ VỤ VIỆC\n\nMức xử lý đề xuất: khiển trách.",
        downstream_handoff=DownstreamHandoff(can_recommend=True, can_create_legal_ticket=True),
    )
    accepted = guard(raw, [rule])
    decision = _decide(_layered(rule), {**_READY, "absence_days": "1 ngày"}, assessment=accepted)
    prose = draft_answer(
        AnswerInput(
            family=ABSENCE,
            facts={**_READY, "absence_days": "1 ngày"},
            fact_lines=["Số ngày nghỉ không phép: 1 ngày"],
            employee_name="An",
            incident_date=date(2026, 3, 1),
            layered=_layered(rule),
            precedent_note="Không có đoạn tiền lệ trong lượt đối chiếu này.",
        ),
        accepted,
    ).prose
    assert accepted.sanction_assessment.recommended_sanction == "WARNING"
    assert accepted.sanction_assessment.status == "SUPPORTED"
    assert accepted.risk_assessment.level == "HIGH"
    assert accepted.downstream_handoff.can_recommend is False
    assert accepted.downstream_handoff.can_create_legal_ticket is False
    assert accepted.downstream_handoff.needs_legal_review is True
    assert accepted.downstream_handoff.hr_decision_required is True
    assert "REJECT_ASSESSMENT" in accepted.downstream_handoff.next_allowed_actions
    assert decision.sanction_type is None
    assert "khiển trách" in decision.rationale
    assert prose.startswith("BIÊN BẢN ĐÁNH GIÁ VỤ VIỆC")


def test_a_law_that_names_one_form_stays_in_the_report_until_an_internal_rule_is_retrieved():
    law = _passage("LAW", "BLLD", "Bộ luật Lao động", "Đi muộn lần đầu thì khiển trách.")
    raw = Assessment(
        case_summary=CaseSummary(violation="đi muộn", facts={"late_days": "1 ngày"}),
        sanction_assessment=SanctionAssessment(
            status="SUPPORTED",
            recommended_sanction="WARNING",
            reason="Luật nêu khiển trách cho đi muộn lần đầu.",
            evidence=[EvidenceRef(chunk_id=str(law.chunk_id), quote=law.content)],
        ),
        risk_assessment=RiskAssessment(level="LOW", reasons=["Theo luật."]),
        downstream_handoff=DownstreamHandoff(can_recommend=True, can_create_legal_ticket=True),
    )
    accepted = guard(raw, [law])
    assert accepted.sanction_assessment.recommended_sanction == "WARNING"
    assert accepted.risk_assessment.level == "HIGH"
    assert accepted.downstream_handoff.can_recommend is False
    assert accepted.downstream_handoff.can_create_legal_ticket is False
    assert accepted.downstream_handoff.legal_basis_found is True
    assert accepted.downstream_handoff.internal_rule_found is False
    assert _decide(_layered(law), assessment=accepted).sanction_type is None


def test_should_synthesize_reads_law_or_internal_rule_and_skips_page_index():
    law = _passage("LAW", "BLLD", "Bộ luật Lao động", _ART_124)
    rule = _passage("INTERNAL_RULE", "NQ", "Nội quy", "Đi muộn lần đầu thì khiển trách.")
    penalty = _passage("DECREE", "283/2026/ND-CP", "Nghị định xử phạt", "Phạt tiền.")
    ready = DecisionInput(family=ABSENCE, facts=_READY, layered=_layered(law))
    assert should_synthesize(ready) is True
    assert should_synthesize(DecisionInput(family=ABSENCE, facts=_READY, layered=_layered(rule))) is True
    assert should_synthesize(DecisionInput(family=ABSENCE, facts=_READY, layered=_layered(penalty))) is False
    assert should_synthesize(DecisionInput(family=ABSENCE, facts=_READY, layered=_layered(law, rule), retrieval="page_index")) is False


@pytest.mark.asyncio
async def test_missing_synthesis_key_does_not_call_the_model(monkeypatch):
    monkeypatch.setattr(settings, "DEEPSEEK_API_KEY", "")

    async def boom(*_args, **_kwargs):
        raise AssertionError("synthesis model was called")

    monkeypatch.setattr("src.features.answer.model.synthesis._complete", boom)
    law = _passage("LAW", "BLLD", "Bộ luật Lao động", _ART_125)
    rule = _passage("INTERNAL_RULE", "NQ", "Nội quy", "Tự ý bỏ việc không phép thì bị sa thải.")
    message = CaseHandoff(
        incident_date=date(2026, 3, 1),
        employee_name="An",
        description="nghỉ 5 ngày không phép",
        family=ABSENCE,
        facts={**_READY, "absence_days": "5 ngày"},
        passages=[_cited(law), _cited(rule)],
        missing_internal_rule=False,
        agent_run_id=uuid4(),
    )
    result = await synthesize(message, _layered(law, rule))
    assert result.sanction_assessment.recommended_sanction is None
    assert result.downstream_handoff.can_recommend is False


@pytest.mark.asyncio
async def test_cite_step_returns_a_handoff_and_does_not_draft_a_sanction():
    law = _passage("LAW", "BLLD-2019", "Bộ luật Lao động", "Kỷ luật lao động sa thải.")
    penalty = _passage("DECREE", "283_2026_ND-CP", "Nghị định xử phạt", "Phạt tiền.")

    async def search(_query: SearchQuery) -> list[Passage]:
        return [law, penalty]

    graph = build_case_graph(search)
    stored = {key: StoredFact(value, True) for key, value in _READY.items()}
    state = await graph.ainvoke(
        {
            "text": "Nghỉ không phép",
            "stored": stored,
            "advance": True,
            "incident_date": date(2026, 3, 1),
            "agent_run_id": uuid4(),
            "description": "nghỉ không phép",
            "employee_name": "An",
        }
    )
    assert "draft" not in state
    assert state["handoff"].missing_internal_rule is True
    assert [item.policy_code for item in state["handoff"].passages] == ["BLLD-2019", "283_2026_ND-CP"]


@pytest.mark.asyncio
async def test_recommend_agent_reads_only_the_handoff_message(monkeypatch):
    async def offline(message, _layered):
        from src.features.answer.model.synthesis import unavailable_assessment

        return unavailable_assessment(message.facts, "offline")

    monkeypatch.setattr("src.features.answer.model.agent.synthesize", offline)
    law = _passage("LAW", "BLLD-2019", "Bộ luật Lao động", "Kỷ luật lao động sa thải.")
    penalty = _passage("DECREE", "283_2026_ND-CP", "Nghị định xử phạt", "Phạt tiền.")
    message = CaseHandoff(
        incident_date=date(2026, 3, 1),
        employee_name="An",
        description="nghỉ không phép",
        family=ABSENCE,
        facts=_READY,
        fact_lines=["Loại hợp đồng: xác định thời hạn"],
        passages=[_cited(law), _cited(penalty)],
        missing_internal_rule=True,
        agent_run_id=uuid4(),
        query="nghỉ không phép",
    )
    state = await build_recommend_agent().ainvoke({"handoff": message})
    assert state["draft"].decision.sanction_type is None
    assert state["draft"].decision.risk_level == "HIGH"
    assert state["layered"].inspection_risk[0].policy_code == "283_2026_ND-CP"


@pytest.mark.asyncio
async def test_publish_stores_a_draft_and_citations_from_the_handoff(monkeypatch):
    async def offline(message, _layered):
        from src.features.answer.model.synthesis import unavailable_assessment

        return unavailable_assessment(message.facts, "offline")

    monkeypatch.setattr("src.features.answer.model.agent.synthesize", offline)
    law = _passage("LAW", "BLLD-2019", "Bộ luật Lao động", "Kỷ luật lao động sa thải.")
    case = Case(id=uuid4(), incident_date=date(2026, 3, 1), status="NEED_INFO")
    handoff = CaseHandoff(
        incident_date=case.incident_date,
        employee_name="An",
        family=ABSENCE,
        facts=_READY,
        passages=[_cited(law)],
        missing_internal_rule=True,
        agent_run_id=uuid4(),
    )
    supersede = AsyncMock()
    insert_recommendation = AsyncMock(side_effect=lambda _session, row: row)
    insert_citations = AsyncMock(return_value=[SimpleNamespace(id=uuid4())])
    monkeypatch.setattr(answers, "supersede_drafts", supersede)
    monkeypatch.setattr(answers, "ensure_violation", AsyncMock(return_value=False))
    monkeypatch.setattr(answers, "insert_recommendation", insert_recommendation)
    monkeypatch.setattr(answers, "insert_citations", insert_citations)

    session = AsyncMock()
    run = SimpleNamespace(run_metadata={})
    session.get.return_value = run
    view = await answers.publish(session, answers.PublishWork(case, handoff, uuid4()))
    assert view.chat_message == run.run_metadata["chat_message"]
    assert "synthesis_assessment" in run.run_metadata

    row = insert_recommendation.await_args.args[1]
    citation = insert_citations.await_args.args[2][0]
    assert row.status == "DRAFT"
    assert row.sanction_type is None
    assert row.risk_level == "HIGH"
    assert row.agent_run_id == view.agent_run_id
    assert citation.chunk_id == law.chunk_id
    assert citation.quote == law.content
    assert case.status == "ANALYZED"
    assert view.legal_basis[0].chunk_id == law.chunk_id
    supersede.assert_awaited_once()


def _cited(passage: Passage) -> CitedPassage:
    return CitedPassage(
        chunk_id=passage.chunk_id,
        policy_version_id=passage.policy_version_id,
        content=passage.content,
        article=passage.article,
        clause=passage.clause,
        point=passage.point,
        score=passage.score,
        policy_type=passage.policy_type,
        policy_code=passage.policy_code,
        policy_name=passage.policy_name,
    )
