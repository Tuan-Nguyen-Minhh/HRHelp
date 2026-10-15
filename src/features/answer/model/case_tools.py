"""MCP-style tools for case documents.

The app does not ship an MCP server yet, so this module keeps the tool schemas
and dispatcher close to the domain code. A future MCP transport can expose
``CASE_TOOL_SCHEMAS`` and call ``call_case_tool`` without duplicating business
rules.
"""

import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any, Literal

from pydantic import BaseModel, ConfigDict, Field

from src.common.clock import utc_now, vietnam_date
from src.common.minutes_profile import MinutesProfile, minutes_profile
from src.common.minutes_text import clean_minutes_text as _clean_minutes_text
from src.features.answer.schema import SANCTION_LABELS, AnswerView, CitationView
from src.features.assessment.schemas import Sections
from src.models.cases import Case

CaseToolName = Literal[
    "create_violation_minutes_form",
    "create_legal_ticket_form",
    "export_case_document_word",
]

_EXPORT_DIR = Path("/tmp/hr_case_exports")
_INCIDENT_FACT_KEYS = {
    "absence_days",
    "late_days",
    "notice_form",
    "violation_time",
    "violation_place",
    "violation_description",
}


class CaseToolCall(BaseModel):
    """One tool request, shaped like an MCP tool invocation."""

    model_config = ConfigDict(str_strip_whitespace=True)

    name: CaseToolName
    arguments: dict[str, Any] = Field(default_factory=dict)


class CaseToolResult(BaseModel):
    """Structured result returned by the case document tools."""

    name: str
    title: str
    content: str
    metadata: dict[str, Any] = Field(default_factory=dict)


@dataclass(frozen=True)
class CaseDocument:
    """A renderable case document."""

    title: str
    sections: tuple[tuple[str, str], ...]
    created_date: date | None = None
    profile: MinutesProfile = field(default_factory=MinutesProfile)
    basis_groups: tuple[str, ...] = ()

    @property
    def content(self) -> str:
        return "\n\n".join([self.title, *[f"{heading}\n{body}" for heading, body in self.sections]])


CASE_TOOL_SCHEMAS: tuple[dict[str, Any], ...] = (
    {
        "name": "create_violation_minutes_form",
        "description": "Create a Vietnamese internal-rule violation minutes form from a case and recommendation.",
        "input_schema": {
            "type": "object",
            "properties": {
                "include_citations": {"type": "boolean", "default": True},
            },
        },
    },
    {
        "name": "create_legal_ticket_form",
        "description": "Create a legal review ticket form from a stored recommendation and drafts.",
        "input_schema": {
            "type": "object",
            "properties": {
                "legal_code": {"type": "string", "description": "Legal officer email or code, when selected."},
            },
        },
    },
    {
        "name": "export_case_document_word",
        "description": "Export a case document to a .docx file.",
        "input_schema": {
            "type": "object",
            "required": ["document_type"],
            "properties": {
                "document_type": {
                    "type": "string",
                    "enum": ["minutes", "ticket", "assessment"],
                },
                "file_name": {"type": "string"},
            },
        },
    },
)


async def call_case_tool(case: Case, answer: AnswerView | None, call: CaseToolCall) -> CaseToolResult:
    """Dispatch one safe case-document tool call."""
    if call.name == "create_violation_minutes_form":
        document = violation_minutes_document(
            case,
            answer,
            include_citations=bool(call.arguments.get("include_citations", True)),
        )
        return CaseToolResult(name=call.name, title=document.title, content=document.content)
    if call.name == "create_legal_ticket_form":
        document = legal_ticket_document(case, answer, legal_code=call.arguments.get("legal_code"))
        return CaseToolResult(name=call.name, title=document.title, content=document.content)
    if call.name == "export_case_document_word":
        document_type = str(call.arguments.get("document_type") or "")
        document = _document_for_type(case, answer, document_type)
        path = export_word_document(document, call.arguments.get("file_name"))
        return CaseToolResult(
            name=call.name,
            title=document.title,
            content=document.content,
            metadata={"format": "docx", "path": str(path), "file_name": path.name},
        )
    raise ValueError(f"Unsupported case tool: {call.name}")


def violation_minutes_document(
    case: Case,
    answer: AnswerView | None,
    *,
    include_citations: bool = True,
) -> CaseDocument:
    """Create a company-rule violation minutes form.

    The structure follows common Vietnamese HR minutes forms: company header,
    time/place, reporter, employee, violation content, damage, employee opinion,
    proposed handling, and signatures. It is not a copy of any third-party form.
    """
    facts = _facts(answer)
    sections = [
        ("Thời gian, địa điểm", _event_line(case)),
        ("Người lập biên bản", "Họ tên: ........................................\nChức vụ/Bộ phận: ................................"),
        ("Người bị lập biên bản", _employee(case)),
        ("Nội dung vi phạm", _violation_content(case, facts)),
        ("Thiệt hại vật chất", facts.get("damage", "Chưa ghi nhận hoặc chưa xác định.")),
        ("Ý kiến của người bị lập biên bản", "................................................................................"),
        ("Hình thức xử lý/đề xuất", _sanction(answer)),
        ("Xác nhận quy trình", _process(answer)),
    ]
    if include_citations and answer is not None:
        sections.append(("Căn cứ tham khảo", _citations(answer)))
    sections.append(("Chữ ký", "Người lập biên bản: ....................\nNgười bị lập biên bản: ...................."))
    return CaseDocument("BIÊN BẢN GHI NHẬN VI PHẠM NỘI QUY", tuple(sections))


def legal_ticket_document(case: Case, answer: AnswerView | None, *, legal_code: str | None = None) -> CaseDocument:
    """Create the legal review ticket form. This does not submit the ticket."""
    if answer is None:
        sections = [
            ("Trạng thái", "Hồ sơ chưa có báo cáo đánh giá để lập phiếu pháp chế."),
            ("Hồ sơ", _case_line(case)),
        ]
        return CaseDocument("PHIẾU CHUYỂN PHÁP CHẾ", tuple(sections))
    sections = [
        ("Hồ sơ", _case_line(case)),
        ("Người lao động", _employee(case)),
        ("Người nhận pháp chế", legal_code or "Chưa chọn."),
        ("Tóm tắt đề xuất", answer.rationale or answer.message),
        ("Mức đề xuất", _sanction(answer)),
        ("Rủi ro", answer.risk_level or "Chưa xác định."),
        ("Quy trình", _process(answer)),
        ("Căn cứ", _citations(answer)),
    ]
    return CaseDocument("PHIẾU CHUYỂN PHÁP CHẾ", tuple(sections))


def assessment_document(
    case: Case,
    answer: AnswerView | None,
    sections: Sections | None = None,
) -> CaseDocument:
    return disciplinary_minutes_document(case, answer, sections)


def disciplinary_minutes_document(
    case: Case,
    answer: AnswerView | None,
    sections: Sections | None = None,
) -> CaseDocument:
    """Create the editable disciplinary minutes form HR reviews before confirming."""
    employee = sections.employee_information if sections is not None else None
    incident = sections.incident_summary if sections is not None else None
    handling = sections.handling_proposal if sections is not None else None
    case_meta = _case_description_meta(case.description or "")
    employee_name = _first(
        None if employee is None else employee.employee_name,
        case.employee_name,
        "........................................",
    )
    employee_code = _first(
        None if employee is None else employee.employee_code,
        case.employee_code,
        "........................................",
    )
    position = _first(None if employee is None else employee.position, case_meta.get("position"), ".....................")
    department = _first(None if employee is None else employee.department, case_meta.get("department"), ".................................")
    created_date = vietnam_date(case.created_at)
    profile = minutes_profile(case.case_code, case.created_at)
    raw_narrative = _first(
        None if incident is None else incident.narrative,
        None if answer is None else answer.rationale,
        case.description,
        "................................................................................",
    )
    if _looks_like_assessment_report(raw_narrative) or _looks_like_case_metadata(raw_narrative):
        raw_narrative = _case_violation_narrative(case, sections)
    narrative_lines = _minutes_bullets(raw_narrative)
    employee_statement = _first(
        None if incident is None else incident.employee_statement,
        "................................................................................",
    )
    result_summary = _assessment_result_summary(sections)
    proposal = _first(
        _handling_level(handling),
        None if sections is not None or handling is None else handling.rationale,
        None if answer is None or not answer.sanction_type else _sanction(answer),
        ".......................................................................................",
    )
    basis_groups = _assessment_basis_groups(answer, sections)
    basis = "; ".join(basis_groups)
    return CaseDocument(
        "BIÊN BẢN XỬ LÝ VI PHẠM KỶ LUẬT",
        (
            (
                "Thông tin người lao động",
                "\n".join(
                    [
                        f"Ông/Bà: {employee_name}",
                        f"Mã số nhân viên: {employee_code}",
                        f"Chức vụ/Chức danh: {position}",
                        f"Phòng ban/Bộ phận: {department}",
                        f"Thời gian lập biên bản: {_minutes_date(created_date)}",
                        f"Địa điểm: {profile.meeting_location}",
                    ]
                ),
            ),
            (
                "I. THÀNH PHẦN THAM DỰ",
                "\n".join(
                    [
                        f"- Đại diện Ban lãnh đạo: {profile.leadership_name}; Chức vụ {profile.leadership_position}",
                        f"- Đại diện nhân sự công ty: {profile.hr_name}; Chức vụ {profile.hr_position}; Phòng ban {profile.hr_department}",
                        f"- Người bị lập biên bản: {employee_name}; Chức vụ {position}; Phòng ban {department}",
                        "- Người làm chứng: Ông/Bà ........................................; Phòng ban ........................................",
                    ]
                ),
            ),
            (
                "II. NỘI DUNG",
                "\n".join(
                    [
                        "- Diễn biến sự việc:",
                        *narrative_lines,
                        f"- Đối chiếu quy định: {basis}",
                        "- Thiệt hại của công ty: ................................................................................ ...............................",
                        f"- Ý kiến của người bị lập biên bản: {employee_statement}",
                        f"- Kết quả đánh giá: {result_summary}",
                        f"- Hình thức xử phạt của Ban lãnh đạo: {proposal}",
                        "- Ý kiến của người làm chứng và đại diện nhân sự công ty:",
                        ".....................................................................................................................................................",
                        ".....................................................................................................................................................",
                    ]
                ),
            ),
        ),
        created_date=created_date,
        profile=profile,
        basis_groups=tuple(basis_groups),
    )


def _assessment_basis_groups(answer: AnswerView | None, sections: Sections | None) -> list[str]:
    if sections is not None and sections.citations:
        groups: dict[str, list[str]] = {}
        for item in _ordered_basis(sections.citations):
            title, *parts = re.split(r"\s+[—–]\s+", item.citation_label or item.title)
            references = groups.setdefault(title, [])
            reference = " — ".join(parts)
            if reference and reference not in references:
                references.append(reference)
        return [f"{title} — {'; '.join(refs)}" if refs else title for title, refs in groups.items()]
    return [_citations(answer)]


def export_word_document(document: CaseDocument, file_name: str | None = None) -> Path:
    """Write one document to .docx and return the local file path."""
    try:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt
    except ImportError as exc:  # pragma: no cover - depends on deployment package set.
        raise RuntimeError("python-docx is required to export Word files.") from exc

    _EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    path = _EXPORT_DIR / _safe_file_name(file_name or document.title)
    if document.title == "BIÊN BẢN XỬ LÝ VI PHẠM KỶ LUẬT":
        _export_disciplinary_minutes(
            document,
            path,
            Document,
            WD_ALIGN_PARAGRAPH,
            WD_TABLE_ALIGNMENT,
            WD_CELL_VERTICAL_ALIGNMENT,
            OxmlElement,
            qn,
            Inches,
            Pt,
        )
        return path
    doc = Document()
    header = doc.add_paragraph("CÔNG TY: ........................................")
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    motto = doc.add_paragraph("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc")
    motto.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = doc.add_heading(document.title, level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for heading, body in document.sections:
        doc.add_heading(heading, level=2)
        for line in body.splitlines() or [""]:
            doc.add_paragraph(line)
    doc.add_paragraph(f"Ngày xuất: {utc_now().date().isoformat()}")
    doc.save(path)
    return path


def _export_disciplinary_minutes(
    document: CaseDocument,
    path: Path,
    document_cls,
    paragraph_alignment,
    table_alignment,
    vertical_alignment,
    oxml_element,
    qn_fn,
    inches,
    pt,
) -> None:
    doc = document_cls()
    section = doc.sections[0]
    section.page_width = inches(210 / 25.4)
    section.page_height = inches(297 / 25.4)
    section.left_margin = section.right_margin = inches(20 / 25.4)
    section.top_margin = section.bottom_margin = inches(20 / 25.4)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn_fn("w:eastAsia"), "Times New Roman")
    normal.font.size = pt(13)
    normal.paragraph_format.space_after = pt(6)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.widow_control = True
    title_style = doc.styles["Title"]
    title_style.paragraph_format.line_spacing = 1.3
    for element in title_style.element.xpath("./w:pPr/w:pBdr | ./w:rPr/w:spacing"):
        element.getparent().remove(element)

    info = _section_map(document)
    profile = document.profile
    _add_header_table(doc, table_alignment, vertical_alignment, oxml_element, qn_fn, pt, document)
    _add_center(doc, "BIÊN BẢN XỬ LÝ VI PHẠM KỶ LUẬT", paragraph_alignment, pt, size=16, bold=True)
    _add_paragraph(doc, "Biên bản xem xét xử lý vi phạm kỷ luật lao động đối với:", pt)
    employee = _lines(info.get("Thông tin người lao động", ""))
    employee_name = _line_value(employee, "Ông/Bà")
    employee_code = _line_value(employee, "Mã số nhân viên")
    position = _line_value(employee, "Chức vụ/Chức danh")
    department = _line_value(employee, "Phòng ban/Bộ phận")
    meeting_time = _line_value(employee, "Thời gian lập biên bản")
    location = _line_value(employee, "Địa điểm")
    _add_two_column_line(
        doc,
        f"Ông/Bà: {employee_name}",
        f"Mã số nhân viên: {employee_code}",
        pt,
        oxml_element,
        qn_fn,
        inches,
    )
    _add_two_column_line(
        doc,
        f"Chức vụ/Chức danh: {position}",
        f"Phòng ban/Bộ phận: {department}",
        pt,
        oxml_element,
        qn_fn,
        inches,
    )
    _add_paragraph(doc, f"Vào lúc: {meeting_time}", pt)
    _add_paragraph(doc, f"Tại: {location}", pt)
    _add_paragraph(doc, "I. THÀNH PHẦN THAM DỰ:", pt, bold=True)
    _add_paragraph(doc, "- Đại diện Ban lãnh đạo:", pt)
    _add_two_column_line(
        doc,
        f"Ông/Bà: {profile.leadership_name}",
        f"Chức vụ: {profile.leadership_position}",
        pt,
        oxml_element,
        qn_fn,
        inches,
    )
    _add_paragraph(doc, "- Đại diện nhân sự công ty:", pt)
    _add_paragraph(doc, f"Ông/Bà: {profile.hr_name}", pt)
    _add_two_column_line(
        doc,
        f"Chức vụ: {profile.hr_position}",
        f"Phòng ban: {profile.hr_department}",
        pt,
        oxml_element,
        qn_fn,
        inches,
    )
    _add_paragraph(doc, "- Người bị lập biên bản:", pt)
    _add_paragraph(doc, f"Ông/Bà: {employee_name or '........................................................................................................................................'}", pt)
    _add_two_column_line(
        doc,
        f"Chức vụ: {position}",
        f"Phòng ban: {department}",
        pt,
        oxml_element,
        qn_fn,
        inches,
    )
    _add_paragraph(doc, "- Người làm chứng:", pt)
    _add_paragraph(doc, "Ông/Bà: .......................................................................................................................................", pt)
    _add_two_column_line(
        doc,
        "Chức vụ: .........................................................",
        "Phòng ban: ...............................................",
        pt,
        oxml_element,
        qn_fn,
        inches,
    )
    content = _lines(info.get("II. NỘI DUNG", ""))
    _add_paragraph(doc, "II. NỘI DUNG:", pt, bold=True)
    for line in content:
        if line.startswith("- Đối chiếu quy định: ") and document.basis_groups:
            _add_paragraph(doc, "- Đối chiếu quy định:", pt)
            for group in document.basis_groups:
                _add_paragraph(doc, f"  - {group.rstrip('.')}.", pt)
        else:
            _add_paragraph(doc, line, pt)
    _add_signature_table(doc, table_alignment, vertical_alignment, oxml_element, qn_fn, pt)
    doc.save(path)


def _add_header_table(doc, table_alignment, vertical_alignment, oxml_element, qn_fn, pt, document: CaseDocument) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = table_alignment.CENTER
    table.autofit = False
    _clear_borders(table, oxml_element, qn_fn)
    left, right = table.rows[0].cells
    left.vertical_alignment = vertical_alignment.TOP
    right.vertical_alignment = vertical_alignment.TOP
    profile = document.profile
    company = profile.company_name.replace(" – CÔNG TY CP", "\n– CÔNG TY CP")
    _cell_text(left, f"{company}\nSố: {profile.document_number}", pt, bold_first=True, center=True)
    left.paragraphs[1].runs[0].bold = True
    _cell_text(
        right,
        f"CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc\n\n{profile.city}, {_minutes_calendar_date(document.created_date)}",
        pt,
        bold_first=True,
        center=True,
    )
    right.paragraphs[1].runs[0].bold = True


def _add_signature_table(doc, table_alignment, vertical_alignment, oxml_element, qn_fn, pt) -> None:
    table = doc.add_table(rows=2, cols=2)
    table.alignment = table_alignment.CENTER
    _clear_borders(table, oxml_element, qn_fn)
    rows = [
        ("Đại diện Ban lãnh đạo\n(Ký tên, ghi họ tên và đóng dấu)\n\n\n", "Người bị lập biên bản\n(Ký tên, ghi rõ họ tên)\n\n\n"),
        ("Đại diện nhân sự công ty\n(Ký tên, ghi rõ họ tên)\n\n\n", "Người làm chứng\n(Ký tên, ghi rõ họ tên)\n\n\n"),
    ]
    for row, values in zip(table.rows, rows, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.vertical_alignment = vertical_alignment.TOP
            _cell_text(cell, value, pt, bold_first=True, center=True)


def _add_two_column_line(doc, left: str, right: str, pt, oxml_element, qn_fn, inches) -> None:
    # Separate cells keep a long value from pushing the next field onto its line.
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    _clear_borders(table, oxml_element, qn_fn)
    section = doc.sections[0]
    column_width = int((section.page_width - section.left_margin - section.right_margin) / 2)
    for index, (cell, value) in enumerate(zip(table.rows[0].cells, (left, right), strict=True)):
        table.columns[index].width = cell.width = column_width
        margin = oxml_element("w:tcMar")
        edge = oxml_element("w:right" if index == 0 else "w:left")
        edge.set(qn_fn("w:w"), "170")  # 3 mm, matching the preview column gap.
        edge.set(qn_fn("w:type"), "dxa")
        margin.append(edge)
        cell._tc.get_or_add_tcPr().append(margin)
        _cell_text(cell, value, pt)


def _clear_borders(table, oxml_element, qn_fn) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = oxml_element("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn_fn(tag))
        if element is None:
            element = oxml_element(tag)
            borders.append(element)
        element.set(qn_fn("w:val"), "nil")
    margins = oxml_element("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        element = oxml_element(f"w:{edge}")
        element.set(qn_fn("w:w"), "0")
        element.set(qn_fn("w:type"), "dxa")
        margins.append(element)
    tbl_pr.append(margins)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(oxml_element("w:cantSplit"))


def _cell_text(cell, text: str, pt, *, bold_first: bool = False, center: bool = False) -> None:
    cell.text = ""
    for index, line in enumerate(text.splitlines() or [""]):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.alignment = 1 if center else 0
        # Literal leaders inside fixed cells also render in Word-compatible viewers.
        line = re.sub(r"\.{10,}", "." * 20, line)
        _minutes_run(paragraph, line, pt, bold=bool(bold_first and index == 0), leader=False)


def _add_center(doc, text: str, paragraph_alignment, pt, *, size: int = 13, bold: bool = False) -> None:
    paragraph = _add_paragraph(doc, text, pt, size=size, bold=bold)
    paragraph.alignment = paragraph_alignment.CENTER


def _add_paragraph(doc, text: str, pt, *, size: int = 13, bold: bool = False):
    paragraph = doc.add_paragraph()
    paragraph.alignment = 3 if not bold and not re.fullmatch(r"[.\s]+", text) else 0
    if bold:
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_before = pt(10 if size == 13 else 14)
        if size == 16:
            paragraph.style = doc.styles["Title"]
            paragraph.paragraph_format.space_after = pt(16)
    elif text.startswith("- Đại diện") or text in {"- Người bị lập biên bản:", "- Người làm chứng:", "- Diễn biến sự việc:", "- Đối chiếu quy định:"}:
        paragraph.paragraph_format.keep_with_next = True
    if text.startswith("  - "):
        text = "• " + text[4:]
        paragraph.paragraph_format.left_indent = pt(19.5)
        paragraph.paragraph_format.first_line_indent = pt(-10)
    section = doc.sections[0]
    width = section.page_width - section.left_margin - section.right_margin
    _minutes_run(paragraph, text, pt, size=size, bold=bold, width=width)
    return paragraph


def _minutes_run(paragraph, text: str, pt, *, size: int = 13, bold: bool = False, width=None, leader: bool = True):
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    if leader and re.search(r"\.{10,}", text):
        text = re.split(r"\.{10,}", text, maxsplit=1)[0].rstrip() + " \t\u00a0"
        paragraph.paragraph_format.tab_stops.add_tab_stop(int(width or pt(460)), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        paragraph.alignment = 0
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    for script in ("ascii", "hAnsi", "eastAsia", "cs"):
        run._element.rPr.rFonts.set(qn(f"w:{script}"), "Times New Roman")
    run.font.size = pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold
    return run


def _section_map(document: CaseDocument) -> dict[str, str]:
    return {heading: body for heading, body in document.sections}


def _lines(value: str) -> list[str]:
    return [line for line in value.splitlines() if line.strip()]


def _line_value(lines: list[str], label: str) -> str:
    prefix = f"{label}:"
    for line in lines:
        if line.startswith(prefix):
            return line[len(prefix) :].strip()
    return "........................................"


def _document_for_type(case: Case, answer: AnswerView | None, document_type: str) -> CaseDocument:
    if document_type == "minutes":
        return violation_minutes_document(case, answer)
    if document_type == "ticket":
        return legal_ticket_document(case, answer)
    if document_type == "assessment":
        return assessment_document(case, answer)
    raise ValueError("document_type must be one of: minutes, ticket, assessment.")


def _event_line(case: Case) -> str:
    date = case.incident_date.isoformat() if case.incident_date else "....../....../......"
    return f"Vào ngày {date}, tại ........................................"


def _minutes_date(value: date | None) -> str:
    return f"..... giờ ..... phút, {_minutes_calendar_date(value)}"


def _minutes_calendar_date(value: date | None) -> str:
    if value is None:
        return "ngày ..... tháng ..... năm ....."
    return f"ngày {value.day:02d} tháng {value.month:02d} năm {value.year}"


def _case_line(case: Case) -> str:
    return f"Mã hồ sơ: {case.case_code}\nNgày vụ việc: {case.incident_date or 'Chưa ghi'}"


def _employee(case: Case) -> str:
    return "\n".join(
        [
            f"Họ tên: {case.employee_name or '........................................'}",
            f"Mã nhân viên: {case.employee_code or '........................................'}",
            "Chức vụ/Bộ phận: ........................................",
        ]
    )


def _facts(answer: AnswerView | None) -> dict[str, str]:
    if answer is None:
        return {}
    standard = answer.standard
    data = {"risk_level": answer.risk_level or ""}
    if standard is not None:
        data["process"] = _process(answer)
    return data


def _violation_content(case: Case, facts: dict[str, str]) -> str:
    lines = [case.description or "................................................................................"]
    if facts.get("risk_level"):
        lines.append(f"Mức rủi ro đánh giá: {facts['risk_level']}")
    return "\n".join(lines)


def _sanction(answer: AnswerView | None) -> str:
    if answer is None or not answer.sanction_type:
        return "Chưa đủ căn cứ để khóa hình thức xử lý."
    return SANCTION_LABELS.get(answer.sanction_type, answer.sanction_type)


def _process(answer: AnswerView | None) -> str:
    if answer is None or not answer.process_steps:
        return "Chưa có dữ liệu quy trình."
    return "\n".join(f"- {step.label}: {'đã xong' if step.done else 'chưa xong'}" for step in answer.process_steps)


def _citations(answer: AnswerView | None) -> str:
    if answer is None:
        return "Chưa có căn cứ."
    rows = [*answer.legal_basis, *answer.internal_rule, *answer.precedent, *answer.inspection_risk]
    if not rows:
        return "Chưa có căn cứ."
    return "\n".join(_citation(item) for item in _ordered_basis(rows)[:8])


def _citation(item: CitationView) -> str:
    location = " ".join(part for part in (item.article, item.clause, item.point) if part)
    label = item.policy_name or item.policy_code or "Văn bản"
    if location:
        label = f"{label} - {location}"
    quote = " ".join(item.quote.split())
    return f"- {label}: {quote}"


def _ordered_basis(items: list[Any]) -> list[Any]:
    return [item for _index, item in sorted(enumerate(items), key=lambda row: (_basis_rank(row[1]), row[0]))]


def _basis_rank(item: Any) -> int:
    policy_type = (getattr(item, "policy_type", None) or "").upper()
    title = " ".join(
        str(value or "").upper()
        for value in (
            getattr(item, "citation_label", None),
            getattr(item, "title", None),
            getattr(item, "policy_name", None),
            getattr(item, "policy_code", None),
        )
    )
    if policy_type == "INTERNAL_RULE" or "NỘI QUY" in title or "NOI QUY" in title or "QUY ĐỊNH" in title:
        return 0
    if policy_type in {"DECREE", "CIRCULAR"} or "NGHỊ ĐỊNH" in title or "NGHI DINH" in title:
        return 1
    if policy_type == "LAW" or "BỘ LUẬT" in title or "BO LUAT" in title or "LUẬT" in title:
        return 2
    return 3


def _minutes_bullets(value: str) -> list[str]:
    clean = _clean_minutes_text(value)
    if not clean:
        return ["  - ................................................................................"]
    lines = [line.strip() for line in clean.splitlines() if line.strip()]
    if len(lines) == 1:
        lines = [part.strip() for part in re.split(r"\s+-\s+|(?<=[.!?])\s+(?=[A-ZÀ-Ỹ0-9])", lines[0]) if part.strip()]
    clean_lines = dict.fromkeys(line.lstrip("-• ").rstrip(".;") for line in lines if line.strip())
    return [f"  - {line}." for line in clean_lines if line and not _looks_like_assessment_report(line)]


def _looks_like_assessment_report(value: str) -> bool:
    upper = value.upper()
    markers = (
        "BIÊN BẢN ĐÁNH GIÁ",
        "BIEN BAN ĐANH GIA",
        "KẾT QUẢ ĐÁNH GIÁ",
        "CĂN CỨ ĐÁNH GIÁ",
        "PHÂN TÍCH ĐIỀU KIỆN",
        "QUY TRÌNH XỬ LÝ",
    )
    return any(marker in upper for marker in markers) or "CHUNK_ID" in upper or "CHUNK ID" in upper


def _looks_like_case_metadata(value: str) -> bool:
    lines = [line.strip().casefold() for line in value.splitlines() if line.strip()]
    return any(line.startswith(("phòng ban:", "chức danh:", "chức vụ:", "vị trí:", "vi phạm:")) for line in lines)


def _case_violation_narrative(case: Case, sections: Sections | None) -> str:
    rows = []
    case_meta = _case_description_meta(case.description or "")
    if case_meta.get("violation"):
        rows.append(case_meta["violation"])
    if case_meta.get("body") and not _looks_like_assessment_report(case_meta["body"]):
        rows.append(case_meta["body"])
    if case.incident_date:
        rows.append(f"Ngày xảy ra: {case.incident_date.isoformat()}")
    if sections is not None:
        for fact in sections.confirmed_facts:
            if fact.fact_key in _INCIDENT_FACT_KEYS and fact.value:
                rows.append(f"{fact.label}: {fact.value}")
    return "\n".join(rows)


def _case_description_meta(description: str) -> dict[str, str]:
    meta: dict[str, str] = {}
    body: list[str] = []
    for line in description.splitlines():
        text = line.strip()
        if not text:
            continue
        key, sep, value = text.partition(":")
        normalized = key.strip().casefold()
        clean_value = value.strip()
        if sep and normalized == "phòng ban":
            meta["department"] = clean_value
        elif sep and normalized in {"chức danh", "chức vụ", "vị trí"}:
            meta["position"] = clean_value
        elif sep and normalized == "vi phạm":
            meta["violation"] = clean_value
        else:
            body.append(text)
    if body:
        meta["body"] = " ".join(body)
    return meta


def _assessment_result_summary(sections: Sections | None) -> str:
    if sections is None:
        return "HR xem xét trên thông tin đã thu thập và căn cứ áp dụng trước khi quyết định hình thức xử lý."
    return _clean_minutes_text(sections.assessment_result.summary) or (
        "HR xem xét trên thông tin đã thu thập và căn cứ áp dụng trước khi quyết định hình thức xử lý."
    )


def _handling_level(handling: Any) -> str | None:
    if handling is None:
        return None
    if not handling.proposed_actions:
        return "Chưa xác định do chưa đủ căn cứ."
    labels = [item.label for item in handling.proposed_actions if item.label]
    return "; ".join(labels) if labels else None


def _first(*values: str | None) -> str:
    for value in values:
        if value is not None and value.strip():
            return value.strip()
    return ""


def _safe_file_name(value: str) -> str:
    raw = value.strip()
    if raw.lower().endswith(".docx"):
        raw = raw[:-5]
    folded = "".join(ch if ch.isalnum() else "_" for ch in raw.lower())
    folded = "_".join(part for part in folded.split("_") if part)
    if not folded:
        folded = "case_document"
    folded = f"{folded}_{utc_now().strftime('%Y%m%d%H%M%S')}"
    return f"{folded}.docx"
